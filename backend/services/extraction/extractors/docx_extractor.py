from __future__ import annotations

import logging
from pathlib import Path
from typing import Optional

from ..base import ExtractedDocument, ExtractionError, SimpleExtractor

try:
    from docx import Document  # type: ignore
except ImportError:  # pragma: no cover - optional dependency
    Document = None

logger = logging.getLogger(__name__)


class DocxExtractor(SimpleExtractor):
    name = "docx"
    supported_mime_types = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    )

    def extract(self, file_path: Path, *, mime_type: Optional[str] = None) -> ExtractedDocument:
        if Document is None:
            raise ExtractionError(
                "python-docx is required for DOCX extraction. Install with `pip install python-docx`."
            )
        logger.info("Extracting DOCX text from %s", file_path)
        doc = Document(str(file_path))
        paragraphs = [para.text for para in doc.paragraphs if para.text]
        text = "\n".join(paragraphs)
        metadata = {
            "extraction_method": "docx_text",
            "content_type": mime_type,
            "paragraph_count": len(paragraphs),
        }
        logger.debug("DOCX extraction produced %d paragraphs for %s", metadata["paragraph_count"], file_path)
        return ExtractedDocument(text=text, metadata=metadata)


__all__ = ["DocxExtractor"]
