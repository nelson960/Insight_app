from __future__ import annotations

import logging
from pathlib import Path
from typing import Optional

from ..base import ExtractedDocument, ExtractionError, SimpleExtractor
from ..normalizer import normalize_text

try:
    from pptx import Presentation  # type: ignore
except ImportError:  # pragma: no cover
    Presentation = None

logger = logging.getLogger(__name__)


class PptxExtractor(SimpleExtractor):
    name = "pptx"
    supported_mime_types = (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/vnd.ms-powerpoint",
    )

    def extract(self, file_path: Path, *, mime_type: Optional[str] = None) -> ExtractedDocument:
        if Presentation is None:
            raise ExtractionError(
                "python-pptx is required for PPTX extraction. Install with `pip install python-pptx`."
            )
        logger.info("Extracting PPTX text from %s", file_path)
        prs = Presentation(str(file_path))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    texts.append(shape.text)
        metadata = {
            "slide_count": len(prs.slides),
            "extraction_method": "pptx_text",
            "content_type": mime_type,
        }
        logger.debug("PPTX extraction processed %d slides for %s", metadata["slide_count"], file_path)
        return ExtractedDocument(text=normalize_text("\n".join(texts)), metadata=metadata)


__all__ = ["PptxExtractor"]
